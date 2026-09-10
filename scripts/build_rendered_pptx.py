#!/usr/bin/env python3
"""Build a layout-safe PowerPoint from browser-rendered HTML slides."""

import json
import re
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
RENDER_DIR = ROOT / ".pptx-render"
METADATA = RENDER_DIR / "metadata.json"
NOTES_MD = ROOT / "Konusmaci-notlari.md"
VIDEO = ROOT / "assets" / "meme.mp4"
OUTPUT = ROOT / "AI-101-for-technology.pptx"
TEMP_OUTPUT = ROOT / ".pptx-render" / "AI-101-for-technology.tmp.pptx"

SLIDE_W = 13.333
SLIDE_H = 7.5

# app.js moves the original sixth slide into the third position.
NOTE_ORDER = [1, 2, 6, 3, 4, 5, 7, 8, 9, 10, 11, 12, 13, 14, 15]


def parse_notes() -> dict[int, str]:
    text = NOTES_MD.read_text(encoding="utf-8")
    pattern = re.compile(
        r"^##\s+(\d+)\.\s+.+?\s+—\s+\d+\s+sn\s*$",
        re.MULTILINE,
    )
    matches = list(pattern.finditer(text))
    sections: dict[int, str] = {}
    for index, match in enumerate(matches):
        start = match.end()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        sections[int(match.group(1))] = text[start:end].strip()
    return sections


def px_to_inches(value: float, full_px: float, full_inches: float) -> float:
    return value / full_px * full_inches


def add_link(slide, rect: dict, href: str, viewport: dict) -> None:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(px_to_inches(rect["x"], viewport["width"], SLIDE_W)),
        Inches(px_to_inches(rect["y"], viewport["height"], SLIDE_H)),
        Inches(px_to_inches(rect["width"], viewport["width"], SLIDE_W)),
        Inches(px_to_inches(rect["height"], viewport["height"], SLIDE_H)),
    )
    shape.fill.background()
    shape.line.fill.background()
    shape.click_action.hyperlink.address = href


def make_video_poster(slide_image: Path, rect: dict, output: Path) -> Path:
    with Image.open(slide_image) as image:
        left = max(0, round(rect["x"]))
        top = max(0, round(rect["y"]))
        right = min(image.width, round(rect["x"] + rect["width"]))
        bottom = min(image.height, round(rect["y"] + rect["height"]))
        image.crop((left, top, right, bottom)).save(output, "PNG")
    return output


def build() -> None:
    data = json.loads(METADATA.read_text(encoding="utf-8"))
    viewport = data["viewport"]
    slides = data["slides"]
    note_sections = parse_notes()

    if len(slides) != 15:
        raise RuntimeError(f"Expected 15 rendered slides, found {len(slides)}")
    if not VIDEO.exists():
        raise FileNotFoundError(f"Missing embedded video: {VIDEO}")

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    prs.core_properties.title = "AI 101 for technology"
    prs.core_properties.author = "Erdinç Yılmaz"
    prs.core_properties.subject = "Digiturk Technology Talks"

    for index, slide_data in enumerate(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        image_path = RENDER_DIR / slide_data["image"]
        if not image_path.exists():
            raise FileNotFoundError(image_path)

        slide.shapes.add_picture(
            str(image_path),
            0,
            0,
            width=Inches(SLIDE_W),
            height=Inches(SLIDE_H),
        )

        for link in slide_data.get("links", []):
            add_link(slide, link["rect"], link["href"], viewport)

        for video_index, video in enumerate(slide_data.get("videos", [])):
            rect = video["rect"]
            poster = make_video_poster(
                image_path,
                rect,
                RENDER_DIR
                / f"video-poster-{index + 1:02d}-{video_index + 1:02d}.png",
            )
            slide.shapes.add_movie(
                str(VIDEO),
                Inches(px_to_inches(rect["x"], viewport["width"], SLIDE_W)),
                Inches(px_to_inches(rect["y"], viewport["height"], SLIDE_H)),
                Inches(px_to_inches(rect["width"], viewport["width"], SLIDE_W)),
                Inches(px_to_inches(rect["height"], viewport["height"], SLIDE_H)),
                poster_frame_image=str(poster),
                mime_type="video/mp4",
            )

        slide.notes_slide.notes_text_frame.text = note_sections[NOTE_ORDER[index]]

    prs.save(TEMP_OUTPUT)
    TEMP_OUTPUT.replace(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
